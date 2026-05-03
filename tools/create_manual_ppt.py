from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


WORKSPACE = Path(r"C:\Users\온유\Workspace\New project 2")
OUTPUT_DIR = WORKSPACE / "artifacts"
OUTPUT_FILE = OUTPUT_DIR / "onyu-page-operation-manual.pptx"


BG = RGBColor(244, 247, 251)
NAVY = RGBColor(24, 33, 53)
BLUE = RGBColor(79, 99, 210)
SKY = RGBColor(234, 241, 255)
GREEN = RGBColor(62, 161, 101)
MINT = RGBColor(232, 248, 239)
ORANGE = RGBColor(255, 138, 61)
PEACH = RGBColor(255, 243, 234)
GRAY = RGBColor(98, 112, 137)
LINE = RGBColor(220, 227, 240)
RED = RGBColor(196, 39, 54)
PINK = RGBColor(255, 241, 244)


def set_slide_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG


def add_title(slide, kicker, title, subtitle):
    tx = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(8.5), Inches(1.4))
    frame = tx.text_frame
    frame.clear()
    frame.word_wrap = True

    p = frame.paragraphs[0]
    r = p.add_run()
    r.text = kicker
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = BLUE

    p2 = frame.add_paragraph()
    r2 = p2.add_run()
    r2.text = title
    r2.font.name = "Malgun Gothic"
    r2.font.size = Pt(26)
    r2.font.bold = True
    r2.font.color.rgb = NAVY

    p3 = frame.add_paragraph()
    r3 = p3.add_run()
    r3.text = subtitle
    r3.font.name = "Malgun Gothic"
    r3.font.size = Pt(12)
    r3.font.color.rgb = GRAY


def add_footer(slide, text="온유네 Page 로컬 운영 매뉴얼"):
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.7), Inches(7.0), Inches(11.9), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()

    tx = slide.shapes.add_textbox(Inches(0.7), Inches(7.05), Inches(4), Inches(0.3))
    p = tx.text_frame.paragraphs[0]
    p.text = text
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(10)
    p.font.color.rgb = GRAY


def add_card(slide, left, top, width, height, title, body_lines, accent, fill):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = LINE
    shape.line.width = Pt(1)

    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, Inches(0.12), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    tx = slide.shapes.add_textbox(left + Inches(0.22), top + Inches(0.18), width - Inches(0.35), height - Inches(0.3))
    frame = tx.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP

    p = frame.paragraphs[0]
    p.text = title
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY

    for line_text in body_lines:
      p2 = frame.add_paragraph()
      p2.text = line_text
      p2.font.name = "Malgun Gothic"
      p2.font.size = Pt(11)
      p2.font.color.rgb = GRAY
      p2.level = 0


def add_bullets(slide, left, top, width, title, items, title_color=NAVY):
    box = slide.shapes.add_textbox(left, top, width, Inches(4.8))
    frame = box.text_frame
    frame.word_wrap = True

    p = frame.paragraphs[0]
    p.text = title
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = title_color

    for item in items:
        p2 = frame.add_paragraph()
        p2.text = f"• {item}"
        p2.font.name = "Malgun Gothic"
        p2.font.size = Pt(14)
        p2.font.color.rgb = NAVY
        p2.space_after = Pt(7)


def add_step(slide, index, left, top, title, lines, accent):
    circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, left, top, Inches(0.55), Inches(0.55))
    circle.fill.solid()
    circle.fill.fore_color.rgb = accent
    circle.line.fill.background()
    t = slide.shapes.add_textbox(left, top + Inches(0.02), Inches(0.55), Inches(0.45))
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(index)
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

    card = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left + Inches(0.7), top - Inches(0.02), Inches(4.9), Inches(1.1)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card.line.color.rgb = LINE

    tx = slide.shapes.add_textbox(left + Inches(0.92), top + Inches(0.1), Inches(4.5), Inches(0.82))
    frame = tx.text_frame
    p = frame.paragraphs[0]
    p.text = title
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY

    for line in lines:
        p2 = frame.add_paragraph()
        p2.text = line
        p2.font.name = "Malgun Gothic"
        p2.font.size = Pt(11)
        p2.font.color.rgb = GRAY


def add_command_block(slide, left, top, width, height, title, commands):
    outer = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    outer.fill.solid()
    outer.fill.fore_color.rgb = RGBColor(255, 255, 255)
    outer.line.color.rgb = LINE

    tag = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left + Inches(0.2), top + Inches(0.16), Inches(1.4), Inches(0.34))
    tag.fill.solid()
    tag.fill.fore_color.rgb = SKY
    tag.line.fill.background()

    tag_text = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.18), Inches(1.4), Inches(0.25))
    p_tag = tag_text.text_frame.paragraphs[0]
    p_tag.alignment = PP_ALIGN.CENTER
    run_tag = p_tag.add_run()
    run_tag.text = title
    run_tag.font.name = "Consolas"
    run_tag.font.size = Pt(10)
    run_tag.font.bold = True
    run_tag.font.color.rgb = BLUE

    code = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.62), width - Inches(0.45), height - Inches(0.78))
    frame = code.text_frame
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.text = commands[0]
    p.font.name = "Consolas"
    p.font.size = Pt(13)
    p.font.color.rgb = NAVY
    for line in commands[1:]:
        p2 = frame.add_paragraph()
        p2.text = line
        p2.font.name = "Consolas"
        p2.font.size = Pt(13)
        p2.font.color.rgb = NAVY


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    banner = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(0.7), Inches(11.9), Inches(5.6))
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(255, 255, 255)
    banner.line.color.rgb = LINE

    add_title(
        slide,
        "LOCAL OPERATION GUIDE",
        "온유네 Page 실행 / 재실행 매뉴얼",
        "PC를 껐다 켠 뒤 다시 어떻게 실행하는지, 어떤 서비스는 계속 켜둬도 되는지 한 번에 정리한 안내서",
    )

    hero = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(7.2), Inches(2.4))
    frame = hero.text_frame
    p = frame.paragraphs[0]
    p.text = "핵심 요약"
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = BLUE
    for line in [
        "Docker는 항상 켜둘 필요는 없지만, PostgreSQL 컨테이너를 쓰려면 Docker Desktop은 켜져 있어야 합니다.",
        "백엔드와 프론트엔드는 개발 서버이므로 PC를 다시 켜면 직접 다시 실행해야 합니다.",
        "Docker Compose의 PostgreSQL은 Docker Desktop만 올라오면 자동 재시작될 가능성이 높습니다. 현재 설정은 restart: unless-stopped 입니다.",
    ]:
        p2 = frame.add_paragraph()
        p2.text = f"• {line}"
        p2.font.name = "Malgun Gothic"
        p2.font.size = Pt(16)
        p2.font.color.rgb = NAVY
        p2.space_after = Pt(8)

    side = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(1.95), Inches(2.8), Inches(3.45))
    side.fill.solid()
    side.fill.fore_color.rgb = SKY
    side.line.fill.background()
    tx = slide.shapes.add_textbox(Inches(9.05), Inches(2.2), Inches(2.35), Inches(2.9))
    tf = tx.text_frame
    p = tf.paragraphs[0]
    p.text = "운영 포인트"
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY
    for line in ["DB = Docker", "API = backend", "화면 = frontend", "브라우저 = localhost:5173"]:
        p2 = tf.add_paragraph()
        p2.text = line
        p2.font.name = "Malgun Gothic"
        p2.font.size = Pt(15)
        p2.font.color.rgb = BLUE
    add_footer(slide)

    # Slide 2
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title(slide, "SERVICE MAP", "무엇을 켜야 하는지 구조 이해", "서비스마다 역할과 켜는 방식이 다릅니다.")
    add_card(slide, Inches(0.8), Inches(1.7), Inches(3.8), Inches(2.1), "Docker Desktop + PostgreSQL", [
        "데이터베이스 역할",
        "항상 켜둘 필요는 없음",
        "백엔드가 DB에 붙을 때 필요",
        "Docker Desktop이 켜져 있으면 컨테이너 자동 재시작 가능",
    ], BLUE, SKY)
    add_card(slide, Inches(4.8), Inches(1.7), Inches(3.8), Inches(2.1), "Backend (Express + Prisma)", [
        "API 서버 역할",
        "PC 재부팅 후 직접 다시 실행 필요",
        "명령: npm.cmd run dev",
        "경로: backend",
    ], GREEN, MINT)
    add_card(slide, Inches(8.8), Inches(1.7), Inches(3.8), Inches(2.1), "Frontend (React + Vite)", [
        "웹 화면 역할",
        "PC 재부팅 후 직접 다시 실행 필요",
        "명령: npm.cmd run dev",
        "경로: frontend",
    ], ORANGE, PEACH)
    add_bullets(slide, Inches(0.95), Inches(4.25), Inches(11.2), "결론", [
        "Docker는 꺼도 됩니다. 다만 다음에 작업할 때 Docker Desktop을 먼저 다시 켜면 됩니다.",
        "백엔드와 프론트엔드는 '개발 서버'라서 재부팅 후 자동으로 안 뜨는 것이 정상입니다.",
        "매번 필요한 것은 보통 3개입니다: Docker Desktop, backend dev server, frontend dev server.",
    ])
    add_footer(slide)

    # Slide 3
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title(slide, "AFTER REBOOT", "PC를 껐다 켠 뒤 다시 실행하는 순서", "가장 추천하는 표준 재실행 순서")
    add_step(slide, 1, Inches(0.9), Inches(1.7), "Docker Desktop 실행", [
        "시작 메뉴에서 Docker Desktop 실행",
        "고래 아이콘이 뜨고 Engine running 상태가 될 때까지 대기",
    ], BLUE)
    add_step(slide, 2, Inches(6.5), Inches(1.7), "PostgreSQL 컨테이너 확인", [
        "자동 시작 안 되었으면 프로젝트 루트에서 docker compose up -d 실행",
        "이미 떠 있으면 그대로 사용",
    ], GREEN)
    add_step(slide, 3, Inches(0.9), Inches(3.2), "백엔드 실행", [
        "backend 폴더로 이동",
        "npm.cmd run dev",
    ], ORANGE)
    add_step(slide, 4, Inches(6.5), Inches(3.2), "프론트엔드 실행", [
        "새 터미널에서 frontend 폴더로 이동",
        "npm.cmd run dev",
    ], BLUE)
    add_step(slide, 5, Inches(0.9), Inches(4.7), "브라우저 확인", [
        "http://localhost:5173 접속",
        "일정 대시보드와 여행 페이지가 보이면 정상",
    ], GREEN)
    add_footer(slide)

    # Slide 4
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title(slide, "COMMANDS", "복붙해서 쓰는 실행 명령어", "필요할 때 그대로 붙여넣으면 되는 명령 모음")
    add_command_block(slide, Inches(0.8), Inches(1.7), Inches(3.9), Inches(3.9), "1. 루트 / Docker", [
        'Set-Location "C:\\Users\\온유\\Workspace\\New project 2"',
        'docker compose up -d',
    ])
    add_command_block(slide, Inches(4.75), Inches(1.7), Inches(3.9), Inches(3.9), "2. 백엔드", [
        'Set-Location "C:\\Users\\온유\\Workspace\\New project 2\\backend"',
        'npm.cmd run dev',
    ])
    add_command_block(slide, Inches(8.7), Inches(1.7), Inches(3.9), Inches(3.9), "3. 프론트엔드", [
        'Set-Location "C:\\Users\\온유\\Workspace\\New project 2\\frontend"',
        'npm.cmd run dev',
    ])
    tip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(5.9), Inches(11.5), Inches(0.7))
    tip.fill.solid()
    tip.fill.fore_color.rgb = SKY
    tip.line.fill.background()
    tx = slide.shapes.add_textbox(Inches(1.1), Inches(6.05), Inches(11.1), Inches(0.35))
    p = tx.text_frame.paragraphs[0]
    p.text = "팁: PowerShell에서 npm이 막히면 npm 대신 npm.cmd를 쓰면 됩니다."
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    add_footer(slide)

    # Slide 5
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title(slide, "STOP OR KEEP RUNNING", "평소에는 어떻게 두면 좋은가", "자주 작업하는지에 따라 다르게 운영하면 됩니다.")
    add_card(slide, Inches(0.8), Inches(1.8), Inches(5.7), Inches(3.4), "계속 작업할 때", [
        "Docker Desktop: 켜둬도 괜찮음",
        "backend / frontend: 작업 중이면 계속 켜두면 편함",
        "노트북 절전이나 재부팅 시 dev server는 다시 실행해야 할 수 있음",
    ], GREEN, MINT)
    add_card(slide, Inches(6.8), Inches(1.8), Inches(5.7), Inches(3.4), "작업을 마칠 때", [
        "Docker Desktop: 꺼도 됨",
        "backend / frontend 터미널: Ctrl + C로 종료 가능",
        "나중에 다시 할 때는 Docker -> backend -> frontend 순서로 재실행",
    ], ORANGE, PEACH)
    add_bullets(slide, Inches(0.95), Inches(5.45), Inches(11.2), "권장 습관", [
        "매일 계속 작업하지 않는다면 Docker를 늘 켜둘 필요는 없습니다.",
        "백엔드와 프론트엔드는 개발 서버라 작업 끝나면 끄는 쪽이 일반적입니다.",
        "DB 데이터는 Docker 볼륨에 저장되므로 컨테이너를 다시 띄워도 보통 유지됩니다.",
    ])
    add_footer(slide)

    # Slide 6
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title(slide, "TROUBLESHOOTING", "자주 생기는 문제와 빠른 해결", "실제 실행 중 많이 마주치는 메시지 기준으로 정리")
    add_card(slide, Inches(0.8), Inches(1.8), Inches(3.75), Inches(2.0), "docker 명령을 못 찾음", [
        "Docker Desktop 설치 확인",
        "새 PowerShell 다시 열기",
        "docker version 으로 확인",
    ], BLUE, SKY)
    add_card(slide, Inches(4.8), Inches(1.8), Inches(3.75), Inches(2.0), "docker engine 연결 실패", [
        "Docker Desktop 실행 여부 확인",
        "Engine running 상태 확인",
        "필요하면 Restart Docker Desktop",
    ], GREEN, MINT)
    add_card(slide, Inches(8.8), Inches(1.8), Inches(3.75), Inches(2.0), "npm 실행 정책 오류", [
        "PowerShell에서는 npm.cmd 사용",
        "예: npm.cmd run dev",
    ], ORANGE, PEACH)
    add_card(slide, Inches(0.8), Inches(4.1), Inches(5.6), Inches(1.8), "Prisma DATABASE_URL 오류", [
        "루트 .env 와 backend/.env 확인",
        "DATABASE_URL 값이 있는지 확인",
    ], RED, PINK)
    add_card(slide, Inches(6.75), Inches(4.1), Inches(5.6), Inches(1.8), "접속 주소 확인", [
        "프론트: http://localhost:5173",
        "백엔드: http://localhost:4000",
    ], BLUE, SKY)
    add_footer(slide)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_FILE))
    return OUTPUT_FILE


if __name__ == "__main__":
    path = build_presentation()
    print(path)
